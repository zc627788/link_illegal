from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from threading import Lock

import numpy as np
import olefile
import pdfplumber
import pypdfium2 as pdfium
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from rapidocr_onnxruntime import RapidOCR
from striprtf.striprtf import rtf_to_text

from .fetch_requests import RequestsFetcher
from .utils import collapse_text_lines, extension_from_url, filename_from_url


WORD_EXTENSIONS = {".doc", ".docx", ".wps", ".rtf"}
TEXT_EXTENSIONS = {".txt", ".html", ".htm", ".xml", ".json", ".csv"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
ZIP_EXTENSIONS = {".zip"}
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ocr_engine = None
_ocr_lock = Lock()


@dataclass(slots=True)
class AttachmentExtraction:
    url: str
    filename: str
    attachment_type: str
    text: str
    text_length: int
    status: str
    error: str | None


def _get_ocr_engine() -> RapidOCR:
    global _ocr_engine
    if _ocr_engine is None:
        with _ocr_lock:
            if _ocr_engine is None:
                _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_image_array(image_array: np.ndarray) -> str:
    engine = _get_ocr_engine()
    result, _ = engine(image_array)
    if not result:
        return ""
    return collapse_text_lines(item[1] for item in result if len(item) >= 2 and item[1])


def _extract_pdf_text(data: bytes, max_pages: int = 0) -> str:
    texts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = pdf.pages if max_pages <= 0 else pdf.pages[:max_pages]
        for page in pages:
            text = page.extract_text() or ""
            if text.strip():
                texts.append(text)
    return collapse_text_lines("\n".join(texts).splitlines())


def _ocr_pdf_text(data: bytes, max_pages: int = 2, scale: float = 2.0) -> str:
    pdf = pdfium.PdfDocument(io.BytesIO(data))
    texts: list[str] = []
    page_count = len(pdf)
    for page_index in range(min(max_pages, page_count)):
        page = pdf[page_index]
        pil_image = page.render(scale=scale).to_pil()
        text = _ocr_image_array(np.array(pil_image))
        if text:
            texts.append(text)
    return collapse_text_lines("\n".join(texts).splitlines())


def _extract_docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return collapse_text_lines(paragraphs)


def _extract_rtf_text(data: bytes) -> str:
    decoded = data.decode("utf-8", errors="ignore")
    return collapse_text_lines(rtf_to_text(decoded).splitlines())


def _extract_html_like_text(data: bytes) -> str:
    decoded = data.decode("utf-8", errors="ignore")
    soup = BeautifulSoup(decoded, "html.parser")
    return collapse_text_lines(soup.get_text("\n").splitlines())


def _extract_plain_text(data: bytes) -> str:
    decoded = data.decode("utf-8", errors="ignore")
    return collapse_text_lines(decoded.splitlines())


def _extract_image_text(data: bytes) -> str:
    image = Image.open(io.BytesIO(data)).convert("RGB")
    return _ocr_image_array(np.array(image))


def _extract_zip_text(data: bytes, max_inner_files: int = 5) -> str:
    texts: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist()[:max_inner_files]:
            lower = name.lower()
            payload = archive.read(name)
            if lower.endswith(".docx"):
                text = _extract_docx_text(payload)
            elif lower.endswith(".rtf"):
                text = _extract_rtf_text(payload)
            elif any(lower.endswith(ext) for ext in TEXT_EXTENSIONS):
                if b"<html" in payload[:500].lower():
                    text = _extract_html_like_text(payload)
                else:
                    text = _extract_plain_text(payload)
            else:
                text = ""
            if text:
                texts.append(f"压缩包成员：{name}")
                texts.append(text)
    return "\n\n".join(texts)


def _looks_like_docx(data: bytes) -> bool:
    return zipfile.is_zipfile(io.BytesIO(data))


def _looks_like_rtf(data: bytes) -> bool:
    prefix = data[:20].lstrip()
    return prefix.startswith(b"{\\rtf")


def _looks_like_pdf(data: bytes) -> bool:
    return data.startswith(b"%PDF")


def _looks_like_image(data: bytes) -> bool:
    prefix = data[:12]
    return (
        prefix.startswith(b"\x89PNG")
        or prefix.startswith(b"\xff\xd8\xff")
        or prefix[:4] == b"RIFF"
        or prefix.startswith((b"II*\x00", b"MM\x00*"))
        or prefix.startswith(b"BM")
    )


def _looks_like_ole_word(data: bytes) -> bool:
    return data.startswith(OLE_MAGIC) and olefile.isOleFile(io.BytesIO(data))


def classify_attachment(url: str, content_type: str, data: bytes | None = None) -> str:
    ext = extension_from_url(url)
    content_type = (content_type or "").lower()
    if ext == ".pdf" or "application/pdf" in content_type or (data and _looks_like_pdf(data)):
        return "pdf"
    if ext in IMAGE_EXTENSIONS or content_type.startswith("image/") or (data and _looks_like_image(data)):
        return "image"
    if ext in ZIP_EXTENSIONS:
        return "zip"
    if ext == ".docx" or (data and _looks_like_docx(data)):
        return "word"
    if ext == ".rtf" or (data and _looks_like_rtf(data)):
        return "rtf"
    if ext in {".doc", ".wps"} or (data and _looks_like_ole_word(data)):
        return "legacy_word"
    if ext in TEXT_EXTENSIONS or content_type.startswith("text/") or "html" in content_type:
        return "text"
    return "other"


def extract_attachment_text(
    fetcher: RequestsFetcher,
    url: str,
    max_bytes: int,
    max_pdf_pages: int,
    enable_ocr: bool,
    ocr_max_pages: int,
    pdf_render_scale: float,
) -> AttachmentExtraction:
    response = fetcher.fetch_binary(url)
    filename = filename_from_url(response.final_url or url)
    attachment_type = classify_attachment(response.final_url or url, response.content_type, response.content or None)
    if response.error:
        return AttachmentExtraction(
            url=url,
            filename=filename,
            attachment_type=attachment_type,
            text="",
            text_length=0,
            status="fetch_error",
            error=response.error,
        )
    if response.status_code != 200:
        return AttachmentExtraction(
            url=url,
            filename=filename,
            attachment_type=attachment_type,
            text="",
            text_length=0,
            status=f"http_{response.status_code}",
            error=None,
        )
    if max_bytes > 0 and len(response.content) > max_bytes:
        return AttachmentExtraction(
            url=url,
            filename=filename,
            attachment_type=attachment_type,
            text="",
            text_length=0,
            status="skipped_too_large",
            error=f"bytes={len(response.content)}",
        )

    try:
        if attachment_type == "pdf":
            text = _extract_pdf_text(response.content, max_pages=max_pdf_pages)
            status = "ok" if text else "ocr_needed"
            if not text and enable_ocr:
                text = _ocr_pdf_text(response.content, max_pages=ocr_max_pages, scale=pdf_render_scale)
                status = "ok_ocr" if text else "ocr_needed"
        elif attachment_type == "word":
            text = _extract_docx_text(response.content)
            status = "ok" if text else "empty_text"
        elif attachment_type == "rtf":
            text = _extract_rtf_text(response.content)
            status = "ok" if text else "empty_text"
        elif attachment_type == "legacy_word":
            return AttachmentExtraction(
                url=url,
                filename=filename,
                attachment_type=attachment_type,
                text="",
                text_length=0,
                status="legacy_word_needs_converter",
                error="legacy .doc/.wps needs external converter",
            )
        elif attachment_type == "text":
            if b"<html" in response.content[:500].lower():
                text = _extract_html_like_text(response.content)
            else:
                text = _extract_plain_text(response.content)
            status = "ok" if text else "empty_text"
        elif attachment_type == "image":
            text = _extract_image_text(response.content) if enable_ocr else ""
            status = "ok_ocr" if text else "image_ocr_needed"
        elif attachment_type == "zip":
            text = _extract_zip_text(response.content)
            status = "ok" if text else "zip_no_supported_members"
        else:
            text = ""
            status = "unsupported_type"

        return AttachmentExtraction(
            url=url,
            filename=filename,
            attachment_type=attachment_type,
            text=text,
            text_length=len(text),
            status=status,
            error=None,
        )
    except Exception as exc:
        return AttachmentExtraction(
            url=url,
            filename=filename,
            attachment_type=attachment_type,
            text="",
            text_length=0,
            status="extract_error",
            error=f"{type(exc).__name__}: {exc}",
        )


def build_attachment_text(attachments: list[AttachmentExtraction]) -> str:
    parts: list[str] = []
    for item in attachments:
        if not item.text:
            continue
        header = f"附件：{item.filename or item.url}"
        parts.append(header)
        parts.append(item.text)
    return "\n\n".join(parts)


def summarize_attachment_statuses(attachments: list[AttachmentExtraction]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for item in attachments:
        counts[item.status] = counts.get(item.status, 0) + 1
    return counts
